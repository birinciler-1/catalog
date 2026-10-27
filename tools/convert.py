#!/usr/bin/env python3
"""
محوّل الإكسل والصور إلى بيانات التطبيق.

المدخلات:
  excel/*.xlsx        ملف لكل شركة (اسم الملف = اسم الشركة، أو أضف عمود "الشركة")
  raw_images/         صور المنتجات، اسم كل صورة = كود المادة (مثال: 01010.png)
                      ويمكن وضعها داخل مجلد باسم الشركة عند تكرار الأكواد بين الشركات

المخرجات:
  data/products.json
  images/thumbs/*.webp   (مصغّرة للشبكة)
  images/full/*.webp     (للتفاصيل والمشاركة)
  reports/missing_images.txt

التشغيل:  python tools/convert.py [--prune] [--force]
"""
import argparse
import colorsys
import datetime
import hashlib
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl
from PIL import Image, ImageOps

ALIASES = {
    "name": ["اسم المادة", "اسم المنتج", "المادة", "name"],
    "code": ["كود المادة", "الكود", "كود", "code"],
    "barcode": ["باركود المادة", "الباركود", "باركود", "barcode"],
    "stock": ["عدد الباقي", "الباقي", "الكمية", "المخزون", "stock"],
    "pack": ["عدد التعبئة", "التعبئة", "pack"],
    "price": ["سعر المفرد", "سعر القطعة", "المفرد", "سعر", "السعر", "price"],
    "carton": ["سعر الكارتون", "سعر الكارتونة", "الكارتون", "carton"],
    "company": ["الشركة", "شركة", "company", "brand"],
    "category": ["القسم", "قسم", "التصنيف", "النوع", "category"],
    "color": ["لون الشركة", "اللون", "لون", "color"],
}

# ألوان بالأسماء (يمكن كتابة اسم اللون أو كود مثل #1976D2 في عمود "لون الشركة")
NAMED_COLORS = {
    "احمر": "#D32F2F", "حمراء": "#D32F2F", "red": "#D32F2F",
    "برتقالي": "#F57C00", "برتقاليه": "#F57C00", "orange": "#F57C00",
    "اصفر": "#F9A825", "صفراء": "#F9A825", "yellow": "#F9A825",
    "اخضر": "#388E3C", "خضراء": "#388E3C", "green": "#388E3C",
    "ازرق": "#1976D2", "زرقاء": "#1976D2", "blue": "#1976D2",
    "سماوي": "#00ACC1", "تركوازي": "#00897B", "cyan": "#00ACC1",
    "بنفسجي": "#7B1FA2", "ارجواني": "#7B1FA2", "purple": "#7B1FA2",
    "وردي": "#E91E63", "زهري": "#E91E63", "pink": "#E91E63",
    "بني": "#6D4C41", "brown": "#6D4C41",
    "رمادي": "#616161", "gray": "#616161", "grey": "#616161",
    "اسود": "#212121", "black": "#212121",
    "ذهبي": "#B8860B", "gold": "#B8860B",
    "كحلي": "#1A237E", "navy": "#1A237E",
    "زيتي": "#689F38", "olive": "#689F38",
}
AUTO_COLORS = ["#F57C00", "#1976D2", "#388E3C", "#D32F2F", "#7B1FA2", "#00ACC1",
               "#F9A825", "#6D4C41", "#E91E63", "#455A64", "#689F38", "#5D4037"]
THEME = ["FFFFFF", "000000", "E7E6E6", "44546A", "4472C4", "ED7D31", "A5A5A5", "FFC000", "5B9BD5", "70AD47"]
IMG_EXT = {".png", ".jpg", ".jpeg", ".webp"}
AR_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789")


def norm_header(s):
    s = unicodedata.normalize("NFKC", str(s or "")).strip().lower()
    s = re.sub(r"[ً-ٟـ]", "", s)
    return re.sub(r"\s+", " ", s)


HEADER_LOOKUP = {norm_header(a): k for k, v in ALIASES.items() for a in v}


def to_text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).translate(AR_DIGITS).strip()


def to_code(v, pad):
    s = to_text(v)
    if isinstance(v, (int, float)) and s.isdigit() and pad and len(s) < pad:
        s = s.zfill(pad)
    return s


def to_num(v, warn, where, scale=1):
    if v is None or v == "":
        return 0
    if isinstance(v, (int, float)):
        if isinstance(v, float) and not v.is_integer():
            warn(f"{where}: القيمة {v} فيها كسور — تم التقريب (إن كانت الأرقام بالآلاف استخدم --thousands)")
        return int(round(v * scale))
    s = str(v).translate(AR_DIGITS)
    s = re.sub(r"[^\d-]", "", s)          # يزيل الفواصل والنقاط والرموز
    return int(s) * scale if s not in ("", "-") else 0


def parse_color(v):
    s = norm_header(v).replace("أ", "ا").replace("إ", "ا").replace("آ", "ا")
    m = re.fullmatch(r"#?([0-9a-f]{6})", s)
    if m:
        return "#" + m.group(1).upper()
    return NAMED_COLORS.get(s)


def apply_tint(hexv, tint):
    r, g, b = (int(hexv[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, sat = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, max(0, min(1, l)), sat)
    return "%02X%02X%02X" % (round(r * 255), round(g * 255), round(b * 255))


def fill_color(cell):
    """لون تعبئة الخلية في الإكسل (إن وُجد ولم يكن أبيض)."""
    f = cell.fill
    if not f or f.fill_type != "solid":
        return None
    c = f.fgColor
    if c.type == "rgb" and isinstance(c.rgb, str) and len(c.rgb) == 8 and c.rgb != "00000000":
        hexv = c.rgb[2:]
    elif c.type == "theme" and c.theme is not None and c.theme < len(THEME):
        hexv = apply_tint(THEME[c.theme], c.tint or 0)
    else:
        return None
    return None if hexv.upper() == "FFFFFF" else "#" + hexv.upper()


def classify(v):
    """يتعرف على نوع العمود من عنوانه (تطابق تام، ثم تطابق جزئي مرن)."""
    h = norm_header(v).replace("أ", "ا").replace("إ", "ا").replace("آ", "ا")
    if not h:
        return None
    if h in HEADER_LOOKUP:
        return HEADER_LOOKUP[h]
    if "كارتون" in h or "carton" in h:
        return "carton"
    if "باركود" in h or "barcode" in h:
        return "barcode"
    if "تعبئ" in h or "pack" in h:
        return "pack"
    if "باقي" in h or "مخزون" in h or "stock" in h:
        return "stock"
    if "شركه" in h or "شركة" in h:
        return "company"
    if "قسم" in h or "تصنيف" in h:
        return "category"
    if "لون" in h:
        return "color"
    if "سعر" in h:
        return "price"
    if h.startswith("كود") or h.endswith("كود") or h == "code":
        return "code"
    if "اسم" in h or "name" in h:
        return "name"
    return None


def find_header(ws):
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=40, values_only=True), start=1):
        cols, shown = {}, {}
        for c_idx, v in enumerate(row):
            k = classify(v)
            if k and k not in cols:
                cols[k] = c_idx
                shown[k] = str(v).strip()
        if "name" in cols and "code" in cols:
            return r_idx, cols, shown
    return None, None, None


def read_excels(excel_dir, pad, thousands, warn, layout):
    items, seen, colors = [], {}, {}
    files = sorted(p for p in excel_dir.glob("*.xlsx") if not p.name.startswith("~$"))
    if not files:
        sys.exit(f"لا توجد ملفات xlsx داخل {excel_dir}")
    for f in files:
        wb = openpyxl.load_workbook(f, data_only=True)
        found = False
        for ws in wb.worksheets:                      # ورقة «الشركات»: اسم الشركة + لونها في مكان واحد
            if norm_header(ws.title) not in ("الشركات", "companies"):
                continue
            hdr = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
            kinds = [classify(v) for v in hdr]
            if "company" in kinds and "color" in kinds:
                ci, ki = kinds.index("company"), kinds.index("color")
                for row in ws.iter_rows(min_row=2, values_only=True):
                    nm = to_text(row[ci]) if ci < len(row) else ""
                    cv = to_text(row[ki]) if ki < len(row) else ""
                    if nm and cv:
                        col = parse_color(cv)
                        if col:
                            colors[nm] = col
                        else:
                            warn(f"{f.name}/{ws.title}: لون غير مفهوم «{cv}» للشركة {nm}")
        for ws in wb.worksheets:
            if norm_header(ws.title) in ("الشركات", "companies", "تعليمات"):
                continue
            hdr_row, cols, shown = find_header(ws)
            if not hdr_row:
                continue
            found = True
            layout.append(f"{f.name} / {ws.title}: " + " | ".join(f"{shown[k]}={k}" for k in sorted(cols, key=cols.get)))
            scale = 1000 if thousands else 1
            empty_run = 0
            for r_idx, cells in enumerate(ws.iter_rows(min_row=hdr_row + 1), start=hdr_row + 1):
                row = [c.value for c in cells]
                get = lambda k: row[cols[k]] if k in cols and cols[k] < len(row) else None
                name, code = to_text(get("name")), to_code(get("code"), pad)
                if not name or not code:
                    empty_run += 1
                    if empty_run > 100:      # ملفات فيها تنسيق لملايين الصفوف الفارغة
                        break
                    continue
                empty_run = 0
                where = f"{f.name}/{ws.title} سطر {r_idx}"
                category = to_text(get("category"))
                company = to_text(get("company")) or category or f.stem
                if "company" not in cols:
                    category = ""
                if company not in colors:
                    col = None
                    if "color" in cols and to_text(get("color")):
                        col = parse_color(to_text(get("color")))
                        if not col:
                            warn(f"{where}: لون غير مفهوم «{to_text(get('color'))}» — استخدم اسم لون أو كود مثل #1976D2")
                    if not col and "company" in cols and cols["company"] < len(cells):
                        col = fill_color(cells[cols["company"]])
                    if col:
                        colors[company] = col
                it = {
                    "code": code, "name": re.sub(r"\s+", " ", name), "company": company,
                    "barcode": re.sub(r"\s+", "", to_text(get("barcode"))),
                    "stock": to_num(get("stock"), warn, where),
                    "pack": to_num(get("pack"), warn, where),
                    "price": to_num(get("price"), warn, where, scale),
                    "carton": to_num(get("carton"), warn, where, scale),
                    "category": category,
                }
                if "carton" not in cols and it["price"] and it["pack"]:      # لا يوجد عمود كارتون: السعر × التعبئة
                    it["carton"] = it["price"] * it["pack"]
                key = (company, code.lower())
                if key in seen:
                    warn(f"{where}: الكود {code} مكرر في {company} — اعتُمد السطر الأخير")
                    items[seen[key]] = it
                else:
                    seen[key] = len(items)
                    items.append(it)
        if not found:
            warn(f"{f.name}: لم أجد صف العناوين (اسم المادة / كود المادة) — تم تجاهل الملف")
    return items, colors


def safe(s):
    return re.sub(r"[^A-Za-z0-9_-]+", "_", s)


def build_image_index(img_dir):
    """يفهرس الصور: (شركة|None, اسم الملف) ثم كل جزء من اسم الملف."""
    exact, tokens = {}, {}
    if not img_dir.exists():
        return exact, tokens
    for p in sorted(img_dir.rglob("*")):
        if p.suffix.lower() not in IMG_EXT or p.name.startswith("."):
            continue
        rel = p.relative_to(img_dir)
        scope = rel.parts[0] if len(rel.parts) > 1 else None
        stem = unicodedata.normalize("NFKC", p.stem).strip().lower()
        exact.setdefault((scope, stem), p)
        for t in re.split(r"[^0-9a-z؀-ۿ]+", stem):
            if t:
                tokens.setdefault((scope, t), []).append(p)
    return exact, tokens


def find_image(it, exact, tokens):
    code = it["code"].lower()
    for scope in (it["company"], None):
        if (scope, code) in exact:
            return exact[(scope, code)]
    for scope in (it["company"], None):          # "1 - 00748.png" أو "00748 (1).png"
        c = tokens.get((scope, code))
        if c and len(c) == 1:
            return c[0]
    return None


def save_webp(src, dst, size, quality, force):
    if dst.exists() and not force and dst.stat().st_mtime >= src.stat().st_mtime:
        return False
    im = ImageOps.exif_transpose(Image.open(src))
    im = im.convert("RGBA" if ("A" in im.getbands() or im.mode == "P") else "RGB")
    im.thumbnail((size, size), Image.LANCZOS)
    dst.parent.mkdir(parents=True, exist_ok=True)
    im.save(dst, "WEBP", quality=quality, method=4)
    return True


def main():
    root = Path(__file__).resolve().parent.parent
    ap = argparse.ArgumentParser()
    ap.add_argument("--excel", default=root / "excel", type=Path)
    ap.add_argument("--images", default=root / "raw_images", type=Path)
    ap.add_argument("--out", default=root, type=Path)
    ap.add_argument("--pad", type=int, default=5, help="تعبئة الأكواد الرقمية القصيرة بأصفار (الافتراضي 5)")
    ap.add_argument("--thousands", action="store_true", help="اضرب الأسعار بـ 1000 (إن كانت مكتوبة 7.75 بدل 7750)")
    ap.add_argument("--prune", action="store_true", help="احذف صور المنتجات المحذوفة من الإكسل")
    ap.add_argument("--force", action="store_true", help="أعد تحويل كل الصور")
    a = ap.parse_args()

    warnings = []
    warn = warnings.append
    layout = []
    items, colors = read_excels(a.excel, a.pad, a.thousands, warn, layout)
    companies = list(dict.fromkeys(i["company"] for i in items))
    auto = []
    for c in companies:                      # شركة بلا لون: يُعطى لوناً تلقائياً غير مكرر
        if c not in colors:
            free = [x for x in AUTO_COLORS if x not in colors.values()] or AUTO_COLORS
            colors[c] = free[len(auto) % len(free)] if free is AUTO_COLORS else free[0]
            auto.append(c)
    colors = {c: colors[c] for c in companies}

    code_cos = {}
    for i in items:
        code_cos.setdefault(i["code"].lower(), set()).add(i["company"])

    exact, tokens = build_image_index(a.images)
    used, missing, converted = set(), [], 0
    for it in items:
        src = find_image(it, exact, tokens)
        if not src:
            it["img"] = ""
            missing.append(f'{it["company"]}\t{it["code"]}\t{it["name"]}')
            continue
        stem = safe(it["code"])
        if len(code_cos[it["code"].lower()]) > 1:
            stem += "__" + hashlib.md5(it["company"].encode()).hexdigest()[:4]
        it["img"] = stem
        used.add(stem)
        t = save_webp(src, a.out / "images/thumbs" / f"{stem}.webp", 360, 80, a.force)
        f = save_webp(src, a.out / "images/full" / f"{stem}.webp", 1000, 85, a.force)
        converted += t or f

    if a.prune:
        for kind in ("thumbs", "full"):
            for p in (a.out / "images" / kind).glob("*.webp"):
                if p.stem not in used:
                    p.unlink()

    payload = {"companies": companies, "colors": colors, "items": items}
    digest = hashlib.sha1(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode()).hexdigest()[:12]
    out = a.out / "data" / "products.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    updated = datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds")
    if out.exists():
        try:
            old = json.loads(out.read_text(encoding="utf-8"))
            if old.get("version") == digest:
                updated = old.get("updated", updated)
        except Exception:
            pass
    out.write_text(json.dumps({"version": digest, "updated": updated, **payload},
                              ensure_ascii=False, separators=(",", ":")), encoding="utf-8")

    rep = a.out / "reports"
    rep.mkdir(exist_ok=True)
    (rep / "columns.txt").write_text("\n".join(layout) + "\n", encoding="utf-8")
    (rep / "missing_images.txt").write_text("الشركة\tالكود\tالاسم\n" + "\n".join(missing) + "\n", encoding="utf-8")

    print(f"المواد: {len(items)} | الشركات: {len(companies)} | صور محوّلة الآن: {converted} | بلا صورة: {len(missing)}")
    for l in layout:
        print("الأعمدة:", l)
    for c in companies[:40]:
        print(f"  {c}: {colors[c]}" + (" (تلقائي)" if c in auto else ""))
    for w in warnings[:30]:
        print("تنبيه:", w)
    if len(warnings) > 30:
        print(f"… و{len(warnings) - 30} تنبيهات أخرى")


if __name__ == "__main__":
    main()
